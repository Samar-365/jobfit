"""File Parsing and Ingestion Service for JobFit.
Extracts clean plain text from PDF, DOCX, and TXT resumes with safety checks.
"""
import os
import uuid
from pathlib import Path
from typing import Tuple
from fastapi import UploadFile, HTTPException
import pypdf
import docx

from app.config import settings

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def validate_file(file: UploadFile) -> str:
    """Validate file extension and return normalized lowercase extension."""
    filename = file.filename or "resume"
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{ext}'. Allowed formats: PDF (.pdf), DOCX (.docx), TXT (.txt)."
        )
    return ext


async def save_uploaded_file(file: UploadFile) -> Tuple[str, Path, int]:
    """
    Safely stream uploaded file into local uploads directory using a UUID filename
    to completely prevent directory traversal and collision.
    Returns: (original_filename, safe_filepath, file_size_bytes)
    """
    ext = validate_file(file)
    safe_filename = f"{uuid.uuid4()}{ext}"
    dest_path = settings.upload_path / safe_filename

    # Prevent path traversal
    if not str(dest_path.resolve()).startswith(str(settings.upload_path.resolve())):
        raise HTTPException(status_code=400, detail="Invalid file destination path.")

    total_bytes = 0
    with open(dest_path, "wb") as buffer:
        while chunk := await file.read(1024 * 64):  # 64KB chunks
            total_bytes += len(chunk)
            if total_bytes > MAX_FILE_SIZE_BYTES:
                buffer.close()
                if dest_path.exists():
                    os.remove(dest_path)
                raise HTTPException(
                    status_code=413,
                    detail=f"File exceeds maximum allowed size of 5MB ({MAX_FILE_SIZE_BYTES} bytes)."
                )
            buffer.write(chunk)

    if total_bytes == 0:
        if dest_path.exists():
            os.remove(dest_path)
        raise HTTPException(status_code=400, detail="Uploaded file is empty (0 bytes).")

    return file.filename or "resume", dest_path, total_bytes


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract readable text from a PDF file using pypdf."""
    text_chunks = []
    try:
        reader = pypdf.PdfReader(str(file_path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail="PDF is password protected or encrypted and cannot be parsed."
                )

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text.strip())

        full_text = "\n\n".join(text_chunks).strip()
        if not full_text:
            raise HTTPException(
                status_code=422,
                detail="Could not extract readable text from PDF. It may be a scanned image-only PDF without OCR."
            )
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF file: {str(e)}")


def extract_text_from_docx(file_path: Path) -> str:
    """Extract readable text from a DOCX file including paragraphs and tables."""
    try:
        doc = docx.Document(str(file_path))
        text_lines = []

        # Extract paragraphs
        for para in doc.paragraphs:
            stripped = para.text.strip()
            if stripped:
                text_lines.append(stripped)

        # Extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_lines.append(" | ".join(row_cells))

        full_text = "\n".join(text_lines).strip()
        if not full_text:
            raise HTTPException(status_code=422, detail="DOCX document contains no readable text.")
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read DOCX file: {str(e)}")


def extract_text_from_txt(file_path: Path) -> str:
    """Extract plain text with encoding fallbacks."""
    raw_bytes = file_path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = raw_bytes.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue

    text = raw_bytes.decode("utf-8", errors="replace").strip()
    if not text:
        raise HTTPException(status_code=422, detail="Text document contains no readable characters.")
    return text


def extract_text_from_file(file_path: Path, file_type: str) -> str:
    """Unified entrypoint for document text extraction."""
    ext = file_type.lower()
    if not ext.startswith("."):
        ext = f".{ext}"

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".txt":
        return extract_text_from_txt(file_path)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type '{ext}'.")
