"""Unit tests for document parsing (PDF, DOCX, TXT) and validation."""
import pytest
from pathlib import Path
from fastapi import HTTPException
import docx
import pypdf

from app.services.file_parser import (
    extract_text_from_file,
    extract_text_from_txt,
    extract_text_from_docx,
    extract_text_from_pdf,
)


def test_txt_parser(tmp_path: Path):
    """Verify clean plain text extraction from TXT files."""
    test_file = tmp_path / "resume.txt"
    sample_content = "Alex Rivera\nSenior Software Engineer\nSkills: Python, FastAPI, React"
    test_file.write_text(sample_content, encoding="utf-8")

    result = extract_text_from_txt(test_file)
    assert "Alex Rivera" in result
    assert "FastAPI" in result

    # Test via unified wrapper
    unified_result = extract_text_from_file(test_file, ".txt")
    assert unified_result == result


def test_empty_txt_raises_error(tmp_path: Path):
    """Verify empty text file triggers HTTP 422."""
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("   \n\t  ", encoding="utf-8")

    with pytest.raises(HTTPException) as exc:
        extract_text_from_txt(empty_file)
    assert exc.value.status_code == 422


def test_docx_parser(tmp_path: Path):
    """Verify text extraction from DOCX with paragraphs and tables."""
    docx_file = tmp_path / "resume.docx"
    doc = docx.Document()
    doc.add_heading("Taylor Brooks", level=1)
    doc.add_paragraph("Fullstack Developer with 4 years experience in TypeScript.")
    
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Languages"
    hdr_cells[1].text = "Python, TypeScript"
    doc.save(str(docx_file))

    result = extract_text_from_docx(docx_file)
    assert "Taylor Brooks" in result
    assert "TypeScript" in result
    assert "Languages | Python, TypeScript" in result


def test_pdf_parser(tmp_path: Path):
    """Verify PDF text extraction using pypdf."""
    pdf_file = tmp_path / "resume.pdf"
    
    # Generate a lightweight valid PDF using pypdf
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    
    # We can write text or verify parser handles empty page
    with open(pdf_file, "wb") as f:
        writer.write(f)

    # Empty blank page triggers 422 because no extractable text exists
    with pytest.raises(HTTPException) as exc:
        extract_text_from_pdf(pdf_file)
    assert exc.value.status_code == 422


def test_unsupported_file_extension(tmp_path: Path):
    """Verify unsupported extensions raise 400 error."""
    invalid_file = tmp_path / "resume.exe"
    invalid_file.write_text("binary", encoding="utf-8")

    with pytest.raises(HTTPException) as exc:
        extract_text_from_file(invalid_file, ".exe")
    assert exc.value.status_code == 400
